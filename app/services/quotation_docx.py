"""Quotation → DOCX (python-docx) — the "editable copy" (plan §8.1).

Same block structure as the PDF renderer, expressed as Word tables:

    logo / tagline
    header table  (Bill To | subject | meta)
    items table   (header row; per item a boxed headline row + an unboxed
                   spec row with a boxed warranty label and the photo)
    totals + note · terms & conditions · bank + signature · footer

Fonts are set BY NAME (Lucida Fax / Calibri / Arial Black / Times New Roman
for ₹) — the admins' Office PCs have them, so the Word copy looks right there
even though the server only ships the open look-alikes used for the PDF.
Dimensions reuse the PDF renderer's sample-point geometry (× S) so the two
documents line up.
"""
from __future__ import annotations

import io
import re
from datetime import date
from typing import Iterable, List, Optional, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..schemas.quotation import NoteStyle, QuotationDraft, QuotationItemIn, RowStyle
from . import quotation_brand as brand
from . import quotation_pdf as P
from .quotation_money import fmt_inr, fmt_qty, fmt_rate
from .quotation_pdf import RenderModel

S = P.S  # sample-pt → page-pt, identical to the PDF

FONT_TITLE = "Arial Black"
FONT_SANS = "Calibri"
FONT_SERIF = "Lucida Fax"
FONT_RUPEE = "Times New Roman"

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xFF, 0, 0)
GREEN = RGBColor(0, 0xB0, 0x50)

THIN = 8    # eighths of a point → 1 pt
THICK = 12  # 1.5 pt

_RED_RE = re.compile(r"\[\[(.+?)\]\]", re.S)


# ----------------------------- low-level helpers ------------------------- #

def _pt(v: float) -> Pt:
    return Pt(round(v * S * 2) / 2)


def _twips(v: float) -> str:
    return str(int(round(v * S * 20)))


def _cell_borders(cell, *, top=None, bottom=None, left=None, right=None) -> None:
    """Set borders on a cell; each edge is an eighths-of-a-point width, 0 for
    none, or None to leave untouched."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        # Word insists on schema order inside tcPr.
        tcPr.insert_element_before(borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                                   "w:tcFitText", "w:vAlign", "w:hideMark")
    for edge, sz in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if sz is None:
            continue
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        if sz == 0:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


def _valign(cell, where: str = "center") -> None:
    cell.vertical_alignment = {"center": WD_ALIGN_VERTICAL.CENTER, "top": WD_ALIGN_VERTICAL.TOP,
                               "bottom": WD_ALIGN_VERTICAL.BOTTOM}[where]


def _table_setup(table, widths_pt: Sequence[float], *, margins=(0.5, 0.5, 1.5, 1.5)) -> None:
    """Fixed layout, exact column widths (sample pt), tight cell margins
    (top, bottom, left, right in sample pt)."""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.insert_element_before(layout, "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    mar = OxmlElement("w:tblCellMar")
    for edge, v in zip(("top", "bottom", "left", "right"), margins):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), _twips(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    for row in table.rows:
        for i, w in enumerate(widths_pt):
            if i < len(row.cells):
                row.cells[i].width = _pt(w)
    # Grid widths too (Word honours tblGrid for fixed layouts).
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_pt):
        gc.set(qn("w:w"), _twips(w))


def _row_height(row, sample_pt: float, exact: bool = False) -> None:
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), _twips(sample_pt))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trPr.append(h)


def _cant_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _header_row(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def _box(table, sz: int) -> None:
    """Outer border around a whole table."""
    rows = table.rows
    for r, row in enumerate(rows):
        for c, cell in enumerate(row.cells):
            _cell_borders(
                cell,
                top=sz if r == 0 else None,
                bottom=sz if r == len(rows) - 1 else None,
                left=sz if c == 0 else None,
                right=sz if c == len(row.cells) - 1 else None,
            )


def _style_run(run, font: str, size: float, *, bold=True, color=BLACK, underline=False) -> None:
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    run.font.size = _pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if underline:
        run.font.underline = True


def _tight(paragraph, *, align=None, space_after=0.0, line=None) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if line is not None:
        pf.line_spacing = _pt(line)
    if align is not None:
        paragraph.alignment = align


def _first_para(cell):
    p = cell.paragraphs[0]
    _tight(p)
    return p


def _add_text(paragraph, text: str, font: str, size: float, *, color=BLACK, red_size: Optional[float] = None,
              bold=True, underline=False) -> None:
    """Append runs for `text`; `[[…]]` becomes red (optionally bigger)."""
    pos = 0
    for m in _RED_RE.finditer(text):
        if m.start() > pos:
            _style_run(paragraph.add_run(text[pos:m.start()]), font, size, color=color, bold=bold, underline=underline)
        _style_run(paragraph.add_run(m.group(1)), font, red_size or size, color=RED, bold=bold, underline=underline)
        pos = m.end()
    if pos < len(text):
        _style_run(paragraph.add_run(text[pos:]), font, size, color=color, bold=bold, underline=underline)


def _para(cell, text: str, font: str, size: float, *, align=WD_ALIGN_PARAGRAPH.LEFT, color=BLACK,
          red_size=None, underline=False, new=False, line=None):
    p = cell.add_paragraph() if new else _first_para(cell)
    # Exact line spacing so a substituted font (Word/LibreOffice without
    # Lucida Fax) cannot inflate row heights.
    _tight(p, align=align, line=line if line is not None else max(size, red_size or 0) * 1.2)
    _add_text(p, text, font, size, color=color, red_size=red_size, underline=underline)
    return p


def _collapse_empty_paras(cell) -> None:
    """Word requires a paragraph after a nested table (python-docx adds one)
    and keeps the cell's original one before it. Both are empty but take a
    full Normal line each — squash them to 1 pt."""
    for para in cell.paragraphs:
        if para.text.strip():
            continue
        _tight(para)
        para.paragraph_format.line_spacing = Pt(1)
        for run in para.runs:
            run.font.size = Pt(1)
        pPr = para._p.get_or_add_pPr()
        rPr = pPr.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            pPr.append(rPr)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "2")  # half-points → 1 pt paragraph mark
        rPr.append(sz)


def _nested(cell, rows: int, cols: int):
    t = cell.add_table(rows=rows, cols=cols)
    _collapse_empty_paras(cell)
    return t


def _picture(paragraph, data: bytes, max_w: float, max_h: float) -> None:
    from PIL import Image as PILImage

    prepared = P._prepare_image_bytes(data)
    with PILImage.open(io.BytesIO(prepared)) as im:
        w, h = P._fit(im.width, im.height, max_w, max_h)
    paragraph.add_run().add_picture(io.BytesIO(prepared), width=_pt(w), height=_pt(h))


def _rupee(cell, amount: str, width: float, size: float = 6.5) -> None:
    """`₹` at the left, amount right-aligned via a right tab stop."""
    p = _first_para(cell)
    _style_run(p.add_run("₹"), FONT_RUPEE, size)
    p.add_run("\t")
    _style_run(p.add_run(amount), FONT_SERIF, size)
    p.paragraph_format.tab_stops.add_tab_stop(_pt(width - 3.5), WD_TAB_ALIGNMENT.RIGHT)


def _page_border(section, sz: int) -> None:
    sectPr = section._sectPr
    pg = OxmlElement("w:pgBorders")
    pg.set(qn("w:offsetFrom"), "text")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "1")
        el.set(qn("w:color"), "000000")
        pg.append(el)
    sectPr.insert_element_before(pg, "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt", "w:vAlign",
                                 "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter",
                                 "w:docGrid", "w:printerSettings", "w:sectPrChange")


# ----------------------------- blocks ------------------------------------ #

def _header_block(doc, d: QuotationDraft, sig: brand.Signatory) -> None:
    t = doc.add_table(rows=2, cols=3)
    _table_setup(t, P.HEADER_COLS, margins=(1.0, 0.5, 1.5, 1.5))
    _row_height(t.rows[0], P.HEADER_MIN_H - P.SUBJECT_ROW_H)
    _row_height(t.rows[1], P.SUBJECT_ROW_H)

    bill = t.cell(0, 0).merge(t.cell(1, 0))
    _para(bill, "QUOTATION", FONT_TITLE, 11.4, underline=True, line=14.0)
    _para(bill, "Bill To,", FONT_TITLE, 6.0, new=True, line=7.5)
    _para(bill, d.customer_name, FONT_SANS, 9.7, new=True, line=11.0)
    for line in d.address_lines:
        _para(bill, line, FONT_SANS, 7.0, new=True, line=8.5)
    if d.customer_gstin:
        _para(bill, f"GSTIN: {d.customer_gstin}", FONT_SANS, 7.0, new=True, line=8.5)
    elif d.customer_pan:
        _para(bill, f"PAN : {d.customer_pan}", FONT_SANS, 7.0, new=True, line=8.5)

    subj = t.cell(1, 1)
    if d.subject_line:
        _para(subj, d.subject_line, FONT_SERIF, 9.7)
        _cell_borders(subj, top=THIN)
    _valign(subj, "center")

    meta = t.cell(0, 2).merge(t.cell(1, 2))
    rows = [
        ("GSTIN", brand.COMPANY_GSTIN), ("Reference No", d.reference or "—"),
        ("Date", d.quotation_date.strftime("%d/%m/%Y")), ("Email", sig.email),
        ("Mobile", sig.mobile), ("WhatsApp No.", brand.COMPANY_WHATSAPP), ("Website", brand.COMPANY_WEBSITE),
    ]
    for i, (k, v) in enumerate(rows):
        p = _first_para(meta) if i == 0 else meta.add_paragraph()
        _tight(p, line=8.7)
        _style_run(p.add_run(k), FONT_SANS, 7.0)
        p.add_run("\t")
        _style_run(p.add_run(f": {v}"), FONT_SANS, 7.0)
        p.paragraph_format.tab_stops.add_tab_stop(_pt(62.0), WD_TAB_ALIGNMENT.LEFT)

    # Borders: outer 1.5, two inner verticals 1.
    _box(t, THICK)
    for r in range(2):
        _cell_borders(t.cell(r, 0), right=THIN)
        _cell_borders(t.cell(r, 1), left=THIN, right=THIN)
        _cell_borders(t.cell(r, 2), left=THIN)


def _items_block(doc, m: RenderModel) -> None:
    d = m.draft
    cols = P.COLS_SL if d.sl_no_visible else P.COLS_NO_SL
    keys = [k for k, _ in cols]
    widths = [w for _, w in cols]
    ix = {k: i for i, k in enumerate(keys)}
    ncol = len(cols)
    compact_layout = all(i.row_style == RowStyle.COMPACT for i in d.items)
    # Slightly tighter than the PDF: leaves headroom for font substitution.
    row_min = P.COMPACT_ROW_MIN_H if compact_layout else P.ROW_A_MIN_H - 2.0

    # Count rows first (python-docx tables are easiest built with a known size).
    n_rows = 1
    for item, img in zip(d.items, m.item_images):
        n_rows += 1
        if item.row_style == RowStyle.DETAILED and (item.spec_lines or item.warranty_label or (item.include_image and img)):
            n_rows += 1
    t = doc.add_table(rows=n_rows, cols=ncol)
    _table_setup(t, widths, margins=(1.0, 1.0, 2.0, 2.0))

    # Header row.
    hdr = t.rows[0]
    _row_height(hdr, P.TABLE_HEADER_H)
    _header_row(hdr)
    for k in keys:
        cell = hdr.cells[ix[k]]
        size = 7.6 if k == "DESC" else 5.6 if k == "SL" else 7.0
        title = P.COL_TITLES[k].replace("&amp;", "&")
        _para(cell, title, FONT_SERIF, size, align=WD_ALIGN_PARAGRAPH.CENTER)
        _valign(cell, "center")
        _cell_borders(cell, top=THICK, bottom=THIN, left=THIN, right=THIN)

    r = 1
    for n, (item, img, lt) in enumerate(zip(d.items, m.item_images, m.totals.line_totals), start=1):
        row = t.rows[r]
        _row_height(row, row_min)
        _cant_split(row)
        cells = row.cells
        for c in cells:
            _valign(c, "center")
            _cell_borders(c, top=THIN, bottom=THIN, left=0, right=0)
        _cell_borders(cells[0], left=THICK)
        _cell_borders(cells[-1], right=THICK)

        if "SL" in ix:
            _para(cells[ix["SL"]], str(n), FONT_SERIF, 6.5, align=WD_ALIGN_PARAGRAPH.CENTER)

        has_brand = bool(item.brand)
        has_model = bool(item.model and item.model.strip())
        if has_brand:
            bc = cells[ix["BRAND"]]
            _para(bc, item.brand, FONT_SERIF, 7.2, align=WD_ALIGN_PARAGRAPH.CENTER)
            if item.brand_sub_label:
                _para(bc, item.brand_sub_label, FONT_SERIF, 5.4, align=WD_ALIGN_PARAGRAPH.CENTER, new=True)
        if has_model:
            mc = cells[ix["MODEL"]]
            lines = [l for l in item.model.split("\n") if l.strip()]
            if len(lines) == 1:
                _para(mc, lines[0], FONT_SERIF, 7.0, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                _para(mc, lines[0], FONT_SERIF, 6.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                for l in lines[1:]:
                    _para(mc, l, FONT_SERIF, 9.7, align=WD_ALIGN_PARAGRAPH.CENTER, new=True)

        if not has_brand and not has_model:
            desc = cells[ix["BRAND"]].merge(cells[ix["DESC"]])
        elif not has_model:
            desc = cells[ix["MODEL"]].merge(cells[ix["DESC"]])
        else:
            desc = cells[ix["DESC"]]
        _para(desc, item.headline, FONT_SERIF, 7.0, align=WD_ALIGN_PARAGRAPH.CENTER, red_size=7.6, line=8.6)

        _rupee(cells[ix["PRICE"]], fmt_inr(item.unit_price), widths[ix["PRICE"]])
        _para(cells[ix["QTY"]], fmt_qty(item.quantity), FONT_SERIF, 6.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _rupee(cells[ix["TOTAL"]], fmt_inr(lt), widths[ix["TOTAL"]])
        r += 1

        show_img = bool(item.include_image and img)
        if item.row_style != RowStyle.DETAILED or not (item.spec_lines or item.warranty_label or show_img):
            continue

        row = t.rows[r]
        _cant_split(row)
        cells = row.cells
        for c in cells:
            _cell_borders(c, top=0, bottom=0, left=0, right=0)
        _cell_borders(cells[0], left=THICK)
        _cell_borders(cells[-1], right=THICK)

        wcell = cells[0].merge(cells[ix["MODEL"]])
        _valign(wcell, "center")
        if item.warranty_label:
            inner = _nested(wcell, 1, 1)
            _table_setup(inner, [sum(widths[: ix["MODEL"] + 1]) - 4.0], margins=(0.5, 0.5, 1.0, 1.0))
            _row_height(inner.rows[0], P.WARRANTY_BOX_H, exact=True)
            ic = inner.cell(0, 0)
            _para(ic, item.warranty_label, FONT_SERIF, 7.0, align=WD_ALIGN_PARAGRAPH.CENTER)
            _valign(ic, "center")
            _cell_borders(ic, top=THIN, bottom=THIN, left=THIN, right=THIN)

        spec = cells[ix["DESC"]]
        first = True
        for line in (item.spec_lines or "").split("\n"):
            if not line.strip():
                continue
            if P._is_all_red(line):
                size, red_size, lead = 7.0, None, 8.3
            elif "[[" in line:
                size, red_size, lead = 4.8, 7.0, 8.3
            else:
                size, red_size, lead = 4.8, None, 6.0
            _para(spec, line, FONT_SERIF, size, red_size=red_size, new=not first, line=lead)
            first = False

        icell = cells[ix["PRICE"]].merge(cells[ix["TOTAL"]])
        _valign(icell, "center")
        if show_img:
            spec_h = 0.0
            if item.spec_lines:
                spec_h = sum(8.3 if ("[[" in l) else 6.0 for l in item.spec_lines.split("\n") if l.strip())
            max_h = P.IMG_MAX_H if not item.spec_lines else max(P.IMG_MIN_H, min(P.IMG_MAX_H, spec_h + 4.0))
            p = _first_para(icell)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent = _pt(30.0)
            _picture(p, img, P.IMG_MAX_W, max_h)
        r += 1


def _totals_block(doc, m: RenderModel) -> None:
    d = m.draft
    gap = 491.0 - P.NOTE_W - sum(P.TOTALS_COLS)
    t = doc.add_table(rows=1, cols=3)
    _table_setup(t, [P.NOTE_W, gap, sum(P.TOTALS_COLS)], margins=(0, 0, 0, 0))
    _row_height(t.rows[0], P.TOTALS_ROW_H * 3)

    note = t.cell(0, 0)
    if d.note_text:
        if d.note_style == NoteStyle.GREEN_ON_BLACK:
            inner = _nested(note, 3, 1)
            _table_setup(inner, [P.NOTE_W], margins=(0, 0, 1.5, 1.5))
            for rr in inner.rows:
                _row_height(rr, P.TOTALS_ROW_H, exact=True)
            mid = inner.cell(1, 0)
            _shade(mid, "000000")
            _valign(mid, "center")
            _para(mid, d.note_text, FONT_SERIF, 7.6, color=GREEN)
        else:
            inner = _nested(note, 1, 1)
            _table_setup(inner, [P.NOTE_W], margins=(1, 1, 3, 3))
            _row_height(inner.rows[0], P.TOTALS_ROW_H * 2, exact=True)
            c = inner.cell(0, 0)
            _valign(c, "center")
            _para(c, d.note_text, FONT_SERIF, 6.7, align=WD_ALIGN_PARAGRAPH.CENTER, color=RED)
            _cell_borders(c, top=THIN, bottom=THIN, left=THIN, right=THIN)

    sub_l, gst_l, grand_l = brand.TOTALS_LABELS[d.totals_label_set.value]
    rows = [(sub_l, m.totals.subtotal), (gst_l.format(rate=fmt_rate(d.gst_rate)), m.totals.gst_amount),
            (grand_l, m.totals.grand_total)]
    tot = _nested(t.cell(0, 2), 3, 2)
    _table_setup(tot, list(P.TOTALS_COLS), margins=(0.5, 0.5, 1.0, 1.0))
    for i, (label, amount) in enumerate(rows):
        _row_height(tot.rows[i], P.TOTALS_ROW_H, exact=True)
        lc, ac = tot.cell(i, 0), tot.cell(i, 1)
        _valign(lc, "center"); _valign(ac, "center")
        _para(lc, label, FONT_SERIF, 6.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        _rupee(ac, fmt_inr(amount), P.TOTALS_COLS[1])
        _cell_borders(lc, top=THIN, bottom=THIN, left=THICK, right=THIN)
        _cell_borders(ac, top=THIN, bottom=THIN, left=THIN, right=THICK)
    _cell_borders(tot.cell(0, 0), top=THICK); _cell_borders(tot.cell(0, 1), top=THICK)
    _cell_borders(tot.cell(2, 0), bottom=THICK); _cell_borders(tot.cell(2, 1), bottom=THICK)


def _terms_block(doc, d: QuotationDraft) -> None:
    terms = d.terms or []
    t = doc.add_table(rows=1 + len(terms), cols=2)
    _table_setup(t, list(P.TC_COLS), margins=(0.3, 0.3, 1.4, 1.0))
    _row_height(t.rows[0], P.TC_HEADER_H - 0.6)
    _para(t.cell(0, 0), "Sl No.", FONT_SERIF, 6.5, align=WD_ALIGN_PARAGRAPH.CENTER, underline=True)
    _para(t.cell(0, 1), "Terms & Conditions", FONT_SERIF, 6.5, underline=True)
    for i, line in enumerate(terms, start=1):
        red = (i - 1) == brand.RED_TERM_INDEX
        _para(t.cell(i, 0), str(i), FONT_SERIF, 6.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        _para(t.cell(i, 1), line, FONT_SERIF, 7.0 if red else 5.4, color=RED if red else BLACK,
              line=7.4 if red else 6.2)
    _box(t, THICK)


def _bank_signature_block(doc, sig: brand.Signatory) -> None:
    t = doc.add_table(rows=1, cols=3)
    _table_setup(t, [P.BANK_W, P.BANK_SIG_GAP, P.SIG_W], margins=(0, 0, 0, 0))

    bank = _nested(t.cell(0, 0), 2, 2)
    kv_w = sum(P.BANK_KV_COLS)
    _table_setup(bank, [kv_w, P.BANK_W - kv_w], margins=(0, 0, 0, 0))
    _row_height(bank.rows[0], P.BAND_H, exact=True)
    band = bank.cell(0, 0).merge(bank.cell(0, 1))
    _shade(band, "000000"); _valign(band, "center")
    _para(band, brand.BANK_BAND_TEXT, FONT_SERIF, 6.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=WHITE)
    kv = bank.cell(1, 0)
    for i, (k, v) in enumerate(brand.BANK_DETAILS):
        p = _first_para(kv) if i == 0 else kv.add_paragraph()
        _tight(p, line=9.8)
        p.paragraph_format.left_indent = _pt(1.2)
        _style_run(p.add_run(k), FONT_SERIF, 6.0)
        p.add_run("\t")
        _style_run(p.add_run(f":  {v}"), FONT_SERIF, 6.0)
        p.paragraph_format.tab_stops.add_tab_stop(_pt(P.BANK_KV_COLS[0]), WD_TAB_ALIGNMENT.LEFT)
    qr = bank.cell(1, 1)
    _valign(qr, "center")
    qp = _first_para(qr)
    qp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _picture(qp, brand.QR_SCAN_PAY.read_bytes(), 70.0, 56.0)
    _box(bank, THIN)

    sigt = _nested(t.cell(0, 2), 5, 1)
    _table_setup(sigt, [P.SIG_W], margins=(0, 0, 0, 0))
    heights = (P.BAND_H, 36.0, 8.6, 6.8, 6.8)
    for row, h in zip(sigt.rows, heights):
        _row_height(row, h, exact=True)
    b = sigt.cell(0, 0)
    _shade(b, "000000"); _valign(b, "center")
    _para(b, brand.SIGNATURE_BAND_TEXT, FONT_SERIF, 6.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=WHITE)
    sp = _first_para(sigt.cell(1, 0)); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _valign(sigt.cell(1, 0), "center")
    _picture(sp, brand.STAMP_SIGNATURE.read_bytes(), 44.0, 37.0)
    for i, (text, size) in enumerate(((sig.name, 7.0), (sig.designation, 5.4), (sig.phones, 5.4)), start=2):
        c = sigt.cell(i, 0)
        _valign(c, "center")
        _para(c, text, FONT_SERIF, size, align=WD_ALIGN_PARAGRAPH.CENTER)
    _box(sigt, THIN)


def _footer_block(doc) -> None:
    t = doc.add_table(rows=2, cols=1)
    _table_setup(t, [491.0], margins=(0.5, 0.5, 1, 1))
    _row_height(t.rows[0], 10.5)
    _row_height(t.rows[1], 12.0)
    _para(t.cell(0, 0), brand.COMPANY_WEBSITE, FONT_SANS, 9.7, align=WD_ALIGN_PARAGRAPH.CENTER, underline=True)
    c = t.cell(1, 0)
    _valign(c, "center")
    _para(c, brand.COMPANY_OFFICE_LINE, FONT_SANS, 9.7, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell_borders(c, top=THICK, bottom=THICK)


def _spacer(doc, sample_pt: float) -> None:
    p = doc.add_paragraph()
    _tight(p)
    p.paragraph_format.line_spacing = _pt(sample_pt)
    _style_run(p.add_run(" "), FONT_SANS, 1.0)


# ----------------------------- entry point ------------------------------- #

def render_quotation_docx(m: RenderModel) -> bytes:
    d = m.draft
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Pt(P.LEFT_MARGIN)
    sec.top_margin = Pt(P.TOP_MARGIN)
    sec.bottom_margin = Pt(P.BOTTOM_MARGIN)
    _page_border(sec, THICK)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_SANS
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    props = doc.core_properties
    props.title = f"Quotation {d.reference or ''} (editable copy)".replace("  ", " ")
    props.author = brand.COMPANY_NAME
    props.subject = "Quotation — editable copy of the issued PDF"

    # Logo + tagline
    p = doc.add_paragraph()
    _tight(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    _picture(p, brand.LOGO_BANNER.read_bytes(), P.LOGO_W, P.LOGO_H)
    tag = doc.add_paragraph()
    _tight(tag, align=WD_ALIGN_PARAGRAPH.CENTER, line=5.6)
    _add_text(tag, brand.COMPANY_TAGLINE, FONT_SERIF, 4.8)

    _header_block(doc, d, m.signatory)
    _items_block(doc, m)
    _spacer(doc, 1.5)
    _totals_block(doc, m)
    _terms_block(doc, d)
    _spacer(doc, 1.0)
    _bank_signature_block(doc, m.signatory)
    _footer_block(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
